"""Parser tests. Fixtures are generated programmatically (no binary files in the repo)."""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
from pathlib import Path

import pytest

from rag_os.domain.documents import ParsedDocument, Section
from rag_os.domain.errors import NotSupported, ParseError
from rag_os.infrastructure.parsers import csv_parser, parser_for, supported_extensions
from rag_os.infrastructure.parsers.csv_parser import CsvParser
from rag_os.infrastructure.parsers.docx import DocxParser
from rag_os.infrastructure.parsers.json_parser import JsonLinesParser, JsonParser, JsonpParser
from rag_os.infrastructure.parsers.pdf import PdfParser
from rag_os.infrastructure.parsers.spreadsheet import XlsParser, XlsxParser
from rag_os.infrastructure.parsers.text import TextParser
from rag_os.infrastructure.parsers.xml_parser import XmlParser
from rag_os.infrastructure.registry import PARSERS

# --------------------------------------------------------------------------- helpers


class NonSeekable(io.RawIOBase):
    """A forward-only stream, like an HTTP/blob download body."""

    def __init__(self, data: bytes) -> None:
        self._src = io.BytesIO(data)

    def readable(self) -> bool:
        return True

    def seekable(self) -> bool:
        return False

    def readinto(self, b) -> int:  # type: ignore[no-untyped-def]
        chunk = self._src.read(min(len(b), 7))  # deliberately short reads
        b[: len(chunk)] = chunk
        return len(chunk)


def sections(doc: ParsedDocument) -> list[Section]:
    return list(doc.sections)


def parse(parser, data: bytes, filename: str) -> tuple[ParsedDocument, list[Section]]:  # type: ignore[no-untyped-def]
    doc = parser.parse(io.BytesIO(data), filename)
    return doc, sections(doc)


def make_pdf(pages: list[list[str]], title: str | None = None) -> bytes:
    """Minimal valid PDF (Helvetica text, correct xref offsets)."""
    objs: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: (
            f"<< /Type /Pages /Kids [{' '.join(f'{4 + 2 * i} 0 R' for i in range(len(pages)))}] "
            f"/Count {len(pages)} >>"
        ).encode(),
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    }
    for i, lines in enumerate(pages):
        ops = ["BT", "/F1 12 Tf", "14 TL", "72 720 Td"]
        for line in lines:
            esc = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({esc}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objs[4 + 2 * i] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>"
        ).encode()
        objs[5 + 2 * i] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
    info = ""
    if title:
        info_id = 4 + 2 * len(pages)
        objs[info_id] = f"<< /Title ({title}) /Author (Ada Lovelace) >>".encode()
        info = f" /Info {info_id} 0 R"
    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for num in sorted(objs):
        offsets[num] = len(out)
        out += f"{num} 0 obj\n".encode() + objs[num] + b"\nendobj\n"
    xref = len(out)
    size = max(objs) + 1
    out += f"xref\n0 {size}\n".encode() + b"0000000000 65535 f \n"
    for num in range(1, size):
        out += f"{offsets[num]:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {size} /Root 1 0 R{info} >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def encrypt_pdf(data: bytes, user_password: str) -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter(clone_from=io.BytesIO(data))
    writer.encrypt(user_password=user_password, owner_password="owner-secret", algorithm="AES-256")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


# --------------------------------------------------------------------------- factory


def test_all_parsers_registered() -> None:
    names = {"pdf", "docx", "xlsx", "xls", "csv", "json", "jsonl", "jsonp", "xml", "text"}
    assert set(PARSERS.names()) >= names
    assert {
        ".pdf",
        ".docx",
        ".xlsx",
        ".xls",
        ".csv",
        ".tsv",
        ".json",
        ".jsonl",
        ".ndjson",
        ".jsonp",
        ".xml",
        ".txt",
        ".md",
        ".log",
    } <= set(supported_extensions())


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("report.PDF", PdfParser),
        ("C:\\docs\\Policy.Docx", DocxParser),
        ("sheets/budget.xlsx", XlsxParser),
        ("legacy.xls", XlsParser),
        ("data.tsv", CsvParser),
        ("feed.ndjson", JsonLinesParser),
        ("api.jsonp", JsonpParser),
        ("config.json", JsonParser),
        ("catalog.xml", XmlParser),
        ("README.md", TextParser),
        ("server.log", TextParser),
    ],
)
def test_factory_picks_by_extension(filename: str, expected: type) -> None:
    assert isinstance(parser_for(filename), expected)


def test_factory_extension_wins_over_content_type() -> None:
    assert isinstance(parser_for("a.csv", "application/pdf"), CsvParser)


def test_factory_falls_back_to_content_type() -> None:
    assert isinstance(parser_for("blob-123", "application/pdf; charset=binary"), PdfParser)
    assert isinstance(parser_for("download", "TEXT/CSV"), CsvParser)


def test_factory_rejects_unknown_extension() -> None:
    with pytest.raises(NotSupported) as exc:
        parser_for("setup.exe")
    assert ".pdf" in str(exc.value)
    assert ".pdf" in exc.value.detail["supported_extensions"]  # type: ignore[operator]


# --------------------------------------------------------------------------- pdf


def test_pdf_pages_title_and_metadata() -> None:
    data = make_pdf(
        [["Employee Handbook", "Welcome aboard (v2)."], [], ["Leave policy: 25 days."]], title="HR Handbook"
    )
    doc, secs = parse(PdfParser(), data, "handbook.pdf")
    assert doc.title == "HR Handbook"
    assert doc.metadata["pages"] == 3
    assert doc.metadata["author"] == "Ada Lovelace"
    assert [s.page for s in secs] == [1, 3]  # blank page 2 yields nothing
    assert "Welcome aboard (v2)." in secs[0].text
    assert "25 days" in secs[1].text
    assert doc.metadata["pages_without_text"] == 1


def test_pdf_title_falls_back_to_first_line_and_streams_non_seekable() -> None:
    data = make_pdf([["Quarterly Results", "Revenue grew."]])
    doc = PdfParser().parse(NonSeekable(data), "q.pdf")  # type: ignore[arg-type]
    assert doc.title == "Quarterly Results"
    assert [s.page for s in doc.sections] == [1]


def test_pdf_sections_are_lazy() -> None:
    doc = PdfParser().parse(io.BytesIO(make_pdf([["a"], ["b"]])), "x.pdf")
    assert not isinstance(doc.sections, list)


def test_pdf_encrypted_with_empty_user_password_is_readable() -> None:
    data = encrypt_pdf(make_pdf([["Open secret text"]]), user_password="")
    doc, secs = parse(PdfParser(), data, "open.pdf")
    assert doc.metadata["encrypted"] is True
    assert "Open secret text" in secs[0].text


def test_pdf_password_protected_raises_parse_error() -> None:
    data = encrypt_pdf(make_pdf([["Top secret"]]), user_password="s3cret")
    with pytest.raises(ParseError, match="password"):
        PdfParser().parse(io.BytesIO(data), "locked.pdf")


def test_pdf_corrupt_raises_parse_error() -> None:
    with pytest.raises(ParseError):
        PdfParser().parse(io.BytesIO(b"this is not a pdf"), "bad.pdf")


# --------------------------------------------------------------------------- docx


def _docx_bytes(with_title: bool = True) -> bytes:
    import docx

    d = docx.Document()
    d.core_properties.title = "Travel Policy" if with_title else ""
    d.core_properties.author = "Policy Team"
    d.add_paragraph("Preamble before any heading.")
    d.add_heading("Scope", level=1)
    d.add_paragraph("Applies to all employees.")
    d.add_heading("Rates", level=2)
    d.add_paragraph("Rates are listed below.")
    table = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("City", "Per diem"), ("Paris", "€90 | incl."), ("Pune", "₹4000")]):
        table.rows[r].cells[0].text = a
        table.rows[r].cells[1].text = b
    d.add_paragraph("Text after the table.")
    d.add_paragraph("Book early", style="List Bullet")
    d.add_heading("Claims", level=1)
    d.add_paragraph("Submit within 30 days.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_docx_document_order_headings_and_tables() -> None:
    doc, secs = parse(DocxParser(), _docx_bytes(), "policy.docx")
    assert doc.title == "Travel Policy"
    assert doc.metadata["author"] == "Policy Team"
    shape = [(s.kind, s.heading_path) for s in secs]
    assert shape == [
        ("text", []),
        ("text", ["Scope"]),
        ("text", ["Scope", "Rates"]),
        ("table", ["Scope", "Rates"]),
        ("text", ["Scope", "Rates"]),
        ("text", ["Claims"]),
    ]
    table = secs[3].text.splitlines()
    assert table[0] == "| City | Per diem |"
    assert table[1] == "| --- | --- |"
    assert table[2] == "| Paris | €90 \\| incl. |"
    assert "- Book early" in secs[4].text


def test_docx_title_falls_back_to_first_heading() -> None:
    doc = DocxParser().parse(io.BytesIO(_docx_bytes(with_title=False)), "p.docx")
    assert doc.title == "Scope"


def test_docx_corrupt_raises_parse_error() -> None:
    with pytest.raises(ParseError):
        DocxParser().parse(io.BytesIO(b"PK\x03\x04garbage"), "bad.docx")


# --------------------------------------------------------------------------- spreadsheets


def _assert_header_repeated(secs: list[Section], header: str) -> None:
    for s in secs:
        assert s.kind == "table"
        assert s.text.splitlines()[0] == header
        assert s.text.splitlines()[1].startswith("| ---")


def test_xlsx_header_repeated_per_block_and_sheets() -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Staff"
    ws.append(["Name", "Joined", "Salary", "Active"])
    for i in range(120):
        ws.append([f"Person {i}", dt.datetime(2024, 1, 1 + i % 28), 1000.0 + i, i % 2 == 0])
        if i == 10:
            ws.append([None, None, None, None])  # fully empty row is skipped
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Key", "Value"])
    ws2.append(["region", "EMEA"])
    wb.properties.title = "HR Data"
    buf = io.BytesIO()
    wb.save(buf)

    doc, secs = parse(XlsxParser(), buf.getvalue(), "staff.xlsx")
    assert doc.title == "HR Data"
    assert doc.metadata["sheets"] == ["Staff", "Notes"]
    staff = [s for s in secs if s.heading_path == ["Staff"]]
    assert len(staff) == 3  # 120 rows / 50 per block
    _assert_header_repeated(staff, "| Name | Joined | Salary | Active |")
    assert "| Person 0 | 2024-01-01 | 1000 | TRUE |" in staff[0].text
    assert sum(len(s.text.splitlines()) - 2 for s in staff) == 120
    notes = [s for s in secs if s.heading_path == ["Notes"]]
    assert notes[0].text.splitlines()[-1] == "| region | EMEA |"
    assert doc.metadata["rows"] == 121


def test_xls_header_repeated_per_block() -> None:
    xlwt = pytest.importorskip("xlwt", reason="xlwt (dev extra) generates the legacy .xls fixture")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Ledger")
    for c, h in enumerate(["Account", "Amount"]):
        ws.write(0, c, h)
    for r in range(1, 76):
        ws.write(r, 0, f"ACC-{r}")
        ws.write(r, 1, float(r))
    buf = io.BytesIO()
    wb.save(buf)

    doc, secs = parse(XlsParser(), buf.getvalue(), "ledger.xls")
    assert doc.metadata["sheets"] == ["Ledger"]
    assert len(secs) == 2
    assert all(s.heading_path == ["Ledger"] for s in secs)
    _assert_header_repeated(secs, "| Account | Amount |")
    assert "| ACC-1 | 1 |" in secs[0].text  # 1.0 rendered as integer
    assert doc.metadata["rows"] == 75


def test_xls_corrupt_raises_parse_error() -> None:
    with pytest.raises(ParseError):
        XlsParser().parse(io.BytesIO(b"\xd0\xcf\x11\xe0 not really ole2"), "bad.xls")


def test_xlsx_corrupt_raises_parse_error() -> None:
    with pytest.raises(ParseError):
        XlsxParser().parse(io.BytesIO(b"nope"), "bad.xlsx")


# --------------------------------------------------------------------------- csv


def test_csv_cp1252_encoding_detected() -> None:
    rows = ["nom,ville,description", "René,Montréal,café crème très apprécié", "Zoé,Québec,élève à l'école"]
    data = ("\r\n".join(rows) + "\r\n").encode("cp1252")
    doc, secs = parse(CsvParser(), data, "clients.csv")
    assert doc.metadata["encoding"] == "cp1252"
    assert "| René | Montréal | café crème très apprécié |" in secs[0].text
    assert "| Zoé | Québec | élève à l'école |" in secs[0].text


def test_csv_header_repeated_and_delimiter_sniffed(tmp_path: Path) -> None:
    path = tmp_path / "orders.csv"
    lines = ["id;product;notes"] + [f'{i};Widget {i};"multi\nline, with; delims"' for i in range(1, 121)]
    path.write_text("\n".join(lines), encoding="utf-8")
    with path.open("rb") as fh:
        doc = CsvParser().parse(fh, path.name)
        secs = sections(doc)
    assert doc.metadata["delimiter"] == ";"
    assert len(secs) == 3
    _assert_header_repeated(secs, "| id | product | notes |")
    assert "| 1 | Widget 1 | multi line, with; delims |" in secs[0].text
    assert doc.metadata["rows"] == 120


def test_tsv_nul_chars_and_empty_rows() -> None:
    data = b"a\tb\n1\x00\t2\n\t\n3\t4\n"
    _, secs = parse(CsvParser(), data, "t.tsv")
    assert secs[0].text.splitlines()[2:] == ["| 1 | 2 |", "| 3 | 4 |"]


def test_csv_utf8_bom_and_non_seekable() -> None:
    data = "\ufeffcol\nvalue ✓\n".encode()
    doc = CsvParser().parse(NonSeekable(data), "x.csv")  # type: ignore[arg-type]
    text = next(iter(doc.sections)).text
    assert text.splitlines()[0] == "| col |"
    assert "value ✓" in text


def test_csv_huge_field_is_guarded(monkeypatch: pytest.MonkeyPatch) -> None:
    original = csv.field_size_limit()
    monkeypatch.setattr(csv_parser, "FIELD_SIZE_LIMIT", 1000)
    data = b'h1,h2\n1,"' + b"x" * 5000 + b"\n2,3\n"  # unbalanced quote swallows the file
    try:
        doc = CsvParser().parse(io.BytesIO(data), "big.csv")
        with pytest.raises(ParseError, match="malformed CSV"):
            sections(doc)
    finally:
        csv.field_size_limit(original)


# --------------------------------------------------------------------------- json family


def test_json_array_records_flattened_and_grouped() -> None:
    records = [
        {
            "id": i,
            "name": f"User {i}",
            "address": {"city": "Pune", "zip": None},
            "tags": ["a", "b"],
            "roles": [{"r": "admin"}],
        }
        for i in range(45)
    ]
    doc, secs = parse(JsonParser(), json.dumps(records).encode(), "users.json")
    assert [s.kind for s in secs] == ["record"] * 3  # 20 + 20 + 5
    first = secs[0].text.split("\n\n")[0].splitlines()
    assert first == [
        "id: 0",
        "name: User 0",
        "address.city: Pune",
        "address.zip: null",
        "tags: a, b",
        "roles[0].r: admin",
    ]
    assert doc.metadata["records"] == 45


def test_json_object_with_record_array() -> None:
    payload = {"title": "Store catalog", "version": 3, "items": [{"sku": "A1"}, {"sku": "B2"}]}
    doc, secs = parse(JsonParser(), json.dumps(payload).encode("utf-8-sig"), "catalog.json")
    assert doc.title == "Store catalog"
    assert secs[0].heading_path == [] and "version: 3" in secs[0].text
    assert secs[1].heading_path == ["items"] and secs[1].text == "sku: A1\n\nsku: B2"


def test_json_invalid_raises_with_position() -> None:
    with pytest.raises(ParseError, match="line 2"):
        JsonParser().parse(io.BytesIO(b'{"a": 1,\n "b": }'), "bad.json")


def test_json_concatenated_documents_fall_back_to_lines() -> None:
    doc, secs = parse(JsonParser(), b'{"a": 1}\n{"a": 2}\n', "x.json")
    assert doc.metadata["format"] == "jsonl"
    assert secs[0].text == "a: 1\n\na: 2"


def test_jsonl_skips_and_counts_bad_lines() -> None:
    data = b'{"n": 1}\nnot json\n\n{"n": 2}\n{broken\n{"n": 3}\n'
    doc, secs = parse(JsonLinesParser(), data, "feed.jsonl")
    assert secs[0].text == "n: 1\n\nn: 2\n\nn: 3"
    assert doc.metadata["skipped_lines"] == 2
    assert doc.metadata["first_skipped_line"] == 2
    assert doc.metadata["records"] == 3


def test_jsonl_leading_bad_lines_counted() -> None:
    doc, secs = parse(JsonLinesParser(), b'oops\n{"n": 1}\n', "feed.ndjson")
    assert doc.metadata["skipped_lines"] == 1
    assert secs[0].text == "n: 1"


def test_jsonl_all_bad_lines_raise_with_line_number() -> None:
    with pytest.raises(ParseError, match="line 2") as exc:
        JsonLinesParser().parse(io.BytesIO(b"\nnope\n{bad\n"), "feed.jsonl")
    assert exc.value.detail == {"line": 2, "bad_lines": 2}


@pytest.mark.parametrize(
    "payload",
    [
        b'cb({"q": "hello", "items": [1, 2]});',
        b'/**/ jQuery3310_123({"q": "hello", "items": [1, 2]})',
        b'/**/ typeof cb === \'function\' && cb({"q": "hello", "items": [1, 2]}); // done (ok)',
        b'  window.app.handlers["x"]({"q": "hello", "items": [1, 2]}) ;\n',
    ],
)
def test_jsonp_callback_stripped(payload: bytes) -> None:
    doc, secs = parse(JsonpParser(), payload, "api.jsonp")
    assert secs[0].text == "q: hello\nitems: 1, 2"
    assert doc.metadata["format"] == "jsonp"


def test_jsonp_array_payload() -> None:
    doc, secs = parse(JsonpParser(), b'/**/ cb([{"a": 1}, {"a": 2}]);', "api.jsonp")
    assert doc.metadata["callback"] == "cb"
    assert secs[0].text == "a: 1\n\na: 2"


@pytest.mark.parametrize(
    "payload",
    [
        b'{"plain": "json, no wrapper"}',
        b"cb({a: 1});",  # JS object literal, not JSON
        b'alert(document.cookie); cb({"a": 1});',  # code outside a single call is never evaluated
        b'cb({"a": 1}',  # unterminated
        b"",
    ],
)
def test_jsonp_malformed_rejected(payload: bytes) -> None:
    with pytest.raises(ParseError):
        JsonpParser().parse(io.BytesIO(payload), "api.jsonp")


# --------------------------------------------------------------------------- xml


def test_xml_repeating_children_become_records() -> None:
    items = "".join(
        f'<ns:item id="{i}"><ns:name>Item &amp; {i}</ns:name><price cur="EUR">{i}.50</price></ns:item>'
        for i in range(30)
    )
    data = f'<?xml version="1.0"?><ns:catalog xmlns:ns="urn:x" version="2">{items}</ns:catalog>'.encode()
    doc, secs = parse(XmlParser(), data, "catalog.xml")
    assert doc.metadata["mode"] == "records"
    assert all(s.kind == "record" and s.heading_path == ["catalog"] for s in secs)
    assert secs[0].text.startswith(
        "catalog[@version=2]\n\nitem[@id=0]\n  name: Item & 0\n  price[@cur=EUR]: 0.50"
    )
    assert len(secs) == 2


def test_xml_document_mode_and_title() -> None:
    data = (
        b"<article><head><title>Safety Rules</title></head>"
        b"<body><p>Wear <b>helmets</b> always.</p></body></article>"
    )
    doc, secs = parse(XmlParser(), data, "rules.xml")
    assert doc.title == "Safety Rules"
    assert doc.metadata["mode"] == "document"
    assert [s.kind for s in secs] == ["text", "text"]
    assert secs[1].text == "body\n  p: Wear always.\n    b: helmets"


def test_xml_root_only() -> None:
    _, secs = parse(XmlParser(), b"<note>Remember the milk</note>", "n.xml")
    assert [s.text for s in secs] == ["note: Remember the milk"]


def test_xml_xxe_rejected() -> None:
    data = b'<?xml version="1.0"?><!DOCTYPE foo [<!ENTITY xxe SYSTEM "file:///etc/passwd">]><foo>&xxe;</foo>'
    with pytest.raises(ParseError, match="unsafe XML"):
        XmlParser().parse(io.BytesIO(data), "xxe.xml")


def test_xml_billion_laughs_rejected() -> None:
    lol = "".join(f'<!ENTITY lol{i} "{f"&lol{i - 1};" * 10}">' for i in range(1, 10))
    data = f'<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol0 "lol">{lol}]><lolz>&lol9;</lolz>'.encode()
    with pytest.raises(ParseError, match="unsafe XML"):
        XmlParser().parse(io.BytesIO(data), "bomb.xml")


def test_xml_malformed_rejected() -> None:
    with pytest.raises(ParseError, match="malformed XML"):
        XmlParser().parse(io.BytesIO(b"<a><b></a>"), "bad.xml")


# --------------------------------------------------------------------------- text


def test_markdown_heading_paths_and_code_fences() -> None:
    md = (
        "# Guide\n\nIntro paragraph.\n\n## Setup\nInstall it.\n\n"
        "```bash\n# not a heading\n\npip install x\n```\n"
        "### Linux\nUse apt.\n## Usage\nRun it.\n"
    )
    doc, secs = parse(TextParser(), md.encode(), "guide.md")
    assert doc.title == "Guide"
    assert [(s.heading_path, s.text) for s in secs] == [
        (["Guide"], "Intro paragraph."),
        (["Guide", "Setup"], "Install it."),
        (["Guide", "Setup"], "```bash\n# not a heading\n\npip install x\n```"),
        (["Guide", "Setup", "Linux"], "Use apt."),
        (["Guide", "Usage"], "Run it."),
    ]


def test_plain_text_paragraphs_encoding_and_nul() -> None:
    data = "Première ligne\x00 du texte.\r\nSuite.\r\n\r\nDeuxième paragraphe très long.\r\n".encode("cp1252")
    doc, secs = parse(TextParser(), data, "notes.txt")
    assert doc.title == "Première ligne du texte."
    assert [s.text for s in secs] == ["Première ligne du texte.\nSuite.", "Deuxième paragraphe très long."]
    assert all(s.heading_path == [] for s in secs)


def test_log_files_skip_heading_detection() -> None:
    doc, secs = parse(TextParser(), b"# not a heading\n2024-01-01 INFO started\n", "app.log")
    assert doc.title is None
    assert secs[0].heading_path == []
    assert secs[0].text.startswith("# not a heading")


def test_text_without_blank_lines_is_bounded() -> None:
    data = ("log line with some words\n" * 3000).encode()
    _, secs = parse(TextParser(), data, "big.log")
    assert len(secs) > 1
    assert all(len(s.text) <= 16_100 for s in secs)


# ---------------------------------------------------------------- front matter
# A document states its own effective date, which the answer prompt uses to prefer the newer of two that
# disagree. Before this, no parser emitted the key ProcessItem reads, so the rule could never fire.


def _parse_text(body: bytes, name: str = "policy.md"):  # type: ignore[no-untyped-def]
    doc = TextParser().parse(io.BytesIO(body), name)
    sections = list(doc.sections)  # metadata is written while the sections stream, so consume first
    return doc, sections


def test_front_matter_supplies_the_effective_date() -> None:
    doc, sections = _parse_text(b"---\neffective_date: 2026-04-01\n---\n# Leave\n\nTwenty-six weeks.\n")
    assert doc.metadata["effective_date"] == "2026-04-01"
    assert doc.title == "Leave", "the fence must not become the title"
    assert sections and sections[0].text == "Twenty-six weeks."


def test_front_matter_is_stripped_from_the_text() -> None:
    """It is metadata. Left in the body it would be embedded, chunked, and quoted back in a citation snippet."""
    _, sections = _parse_text(b"---\neffective_date: 2026-04-01\n---\n# Leave\n\nBody.\n")
    joined = "\n".join(s.text for s in sections)
    assert "effective_date" not in joined and "---" not in joined


@pytest.mark.parametrize("value", [b"1 April 2026", b"2026/04/01", b"tomorrow", b"", b"2026-13-45x"])
def test_a_date_that_is_not_iso_is_ignored(value: bytes) -> None:
    """The field is sortable in the index and is rendered into the prompt, so a value that cannot be ordered
    is worse than none: the model would be asked to prefer the most recent of two things it cannot compare."""
    doc, _ = _parse_text(b"---\neffective_date: " + value + b"\n---\n# T\n\nBody.\n")
    assert "effective_date" not in doc.metadata


def test_only_known_keys_are_taken_from_a_document() -> None:
    """Front matter is untrusted input, and it lands in the same dict as the parser's own stats - so a
    document could otherwise overwrite `lines` or `encoding`, or inject anything it liked."""
    doc, _ = _parse_text(b"---\neffective_date: 2026-04-01\nowner: HR\nlines: 99999\n---\n# T\n\nBody.\n")
    assert doc.metadata["effective_date"] == "2026-04-01"
    assert "owner" not in doc.metadata
    assert doc.metadata["lines"] != 99999, "a document must not be able to rewrite parser statistics"


def test_a_document_without_front_matter_is_unchanged() -> None:
    """The overwhelmingly common case, and the one a regression here would break silently."""
    doc, sections = _parse_text(b"# Password Policy\n\nMinimum length is 14 characters.\n")
    assert doc.title == "Password Policy" and "effective_date" not in doc.metadata
    assert [s.text for s in sections] == ["Minimum length is 14 characters."]


def test_an_unterminated_fence_does_not_swallow_the_document() -> None:
    """`---` is also a Markdown horizontal rule, so an opening fence with no closing one is a real shape. It
    must cost at most the scan window, never the body."""
    body = b"---\n" + b"\n".join(b"line %d" % i for i in range(60)) + b"\n"
    doc, sections = _parse_text(body)
    assert "effective_date" not in doc.metadata
    assert sections, "the document must still produce content"
    assert "line 59" in "\n".join(s.text for s in sections)

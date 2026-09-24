"""Spreadsheet parsers: .xlsx (openpyxl, read-only streaming) and legacy .xls (xlrd).

Each sheet becomes header-repeating row blocks (kind="table", heading_path=[sheet name]).
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import IO, Any

import openpyxl
import xlrd

from rag_os.application.ports import DocumentParser
from rag_os.domain.documents import ParsedDocument, Section
from rag_os.domain.errors import ParseError
from rag_os.infrastructure.parsers._common import ensure_seekable, iso, table_sections
from rag_os.infrastructure.registry import PARSERS

_ZIP_MAGIC = b"PK\x03\x04"


@PARSERS.register("xlsx", description="Excel .xlsx/.xlsm via openpyxl (read-only); header-repeating rows.")
class XlsxParser(DocumentParser):
    extensions = (".xlsx", ".xlsm")
    content_types = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel.sheet.macroenabled.12",
    )

    def parse(self, stream: IO[bytes], filename: str) -> ParsedDocument:
        try:
            wb = openpyxl.load_workbook(ensure_seekable(stream), read_only=True, data_only=True)
        except Exception as e:
            raise ParseError(f"cannot read Excel workbook '{filename}': {e}") from e
        props = wb.properties
        meta: dict[str, Any] = {"sheets": list(wb.sheetnames), "rows": 0}
        meta.update(
            {
                k: v
                for k, v in {
                    "author": props.creator,
                    "created": iso(props.created),
                    "modified": iso(props.modified),
                }.items()
                if v
            }
        )
        doc = ParsedDocument(title=(props.title or "").strip() or None, metadata=meta)
        doc.sections = self._sections(wb, doc.metadata)
        return doc

    @staticmethod
    def _sections(wb: Any, stats: dict[str, Any]) -> Iterator[Section]:
        try:
            for ws in wb.worksheets:
                ws.reset_dimensions()  # declared dimensions are often wrong; read every row
                yield from table_sections(ws.iter_rows(values_only=True), [ws.title], stats)
        except ParseError:
            raise
        except Exception as e:
            raise ParseError(f"corrupt worksheet data: {e}") from e
        finally:
            wb.close()


@PARSERS.register("xls", description="Legacy Excel .xls via xlrd; header-repeating row blocks.")
class XlsParser(DocumentParser):
    extensions = (".xls",)
    content_types = ("application/vnd.ms-excel",)

    def parse(self, stream: IO[bytes], filename: str) -> ParsedDocument:
        data = stream.read() or b""
        if data.startswith(_ZIP_MAGIC):  # an .xlsx saved with a .xls name
            return XlsxParser().parse(io.BytesIO(data), filename)
        try:
            book = xlrd.open_workbook(file_contents=data, on_demand=True)
        except Exception as e:
            raise ParseError(f"cannot read legacy Excel workbook '{filename}': {e}") from e
        meta: dict[str, Any] = {"sheets": book.sheet_names(), "rows": 0}
        if book.user_name:
            meta["author"] = book.user_name
        doc = ParsedDocument(title=None, metadata=meta)
        doc.sections = self._sections(book, doc.metadata)
        return doc

    @staticmethod
    def _sections(book: Any, stats: dict[str, Any]) -> Iterator[Section]:
        try:
            for index in range(book.nsheets):
                sheet = book.sheet_by_index(index)
                rows = (
                    [_xls_value(cell, book.datemode) for cell in sheet.row(r)] for r in range(sheet.nrows)
                )
                yield from table_sections(rows, [sheet.name], stats)
                book.unload_sheet(index)
        except ParseError:
            raise
        except Exception as e:
            raise ParseError(f"corrupt worksheet data: {e}") from e
        finally:
            book.release_resources()


def _xls_value(cell: Any, datemode: int) -> object:
    ctype = cell.ctype
    if ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, datemode)
        except Exception:
            return cell.value
    if ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value, "#ERR")
    if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    return cell.value
